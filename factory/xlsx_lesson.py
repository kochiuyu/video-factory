#!/usr/bin/env python3
"""YAML stays the source. This file is the faculty/RA edit sheet.

Export a lesson to xlsx. Edit speech and on-screen labels. Import merges
those cells back, writes a .bak once, and refuses the write if validate fails.

  python3 factory/xlsx_lesson.py export lessons/STAT/STAT-M01-DE.yaml
  python3 factory/xlsx_lesson.py import lessons/STAT/STAT-DE-01.yaml sheet.xlsx
"""
from __future__ import annotations

import sys
from pathlib import Path

_FACTORY = Path(__file__).resolve().parent
sys.path.insert(0, str(_FACTORY))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from validate import load_lesson, validate

LANGS = ("en", "zh", "zh_hk")
TRIO_STEMS = (
    "kicker", "title",
    "left", "right",
    "col1", "col2", "col3",
    "step1", "step2", "step3",
    "setup", "result", "q1", "q2",
    "formula", "gloss", "footer", "note", "unit",
    "row1", "row2", "row3",
    "item1", "item2", "item3", "item4",
    "claim", "next",
    "c11", "c12", "c13", "c21", "c22", "c23", "c31", "c32", "c33",
)
SCALAR = ("highlight", "rows", "cols", "chart", "number")
SPEECH = ("en", "zh", "zh_hk", "s1_en", "s1_zh", "s1_zh_hk", "s2_en", "s2_zh", "s2_zh_hk", "s3_en", "s3_zh", "s3_zh_hk")
SKIP_IMPORT = {"beat_id", "words_en", "chars_zh"}

NAVY = "1B3A4B"
GOLD = "C4A35A"
YELLOW = "FFF4CC"
GREY = "F0EBE3"
WHITE = "FFFFFF"
INK = "1A1A1A"

VISUALS = (
    "title_card", "definition_card", "two_column", "formula", "takeaway",
    "example", "three_column", "big_number", "steps", "module_open",
    "module_close", "payoff_matrix", "table_grid", "schematic_chart", "dag",
)


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _beat_columns(beats: list[dict]) -> list[str]:
    present = set()
    for beat in beats:
        present.update(beat.keys())
    cols = ["beat_id", "visual", "seconds"]
    for stem in ("kicker", "title"):
        cols.extend(f"{stem}_{lang}" for lang in LANGS)
    cols.extend(SPEECH)
    cols.extend(["words_en", "chars_zh"])
    for stem in TRIO_STEMS:
        if stem in {"kicker", "title"}:
            continue
        if any(f"{stem}_{lang}" in present or stem in present for lang in LANGS):
            if stem in present and not any(f"{stem}_{lang}" in present for lang in LANGS):
                cols.append(stem)
            else:
                cols.extend(f"{stem}_{lang}" for lang in LANGS)
    for stem in SCALAR:
        if stem in present:
            cols.append(stem)
    return cols


def _cell_value(beat: dict, col: str):
    if col == "beat_id":
        return beat.get("id") or ""
    if col in {"words_en", "chars_zh"}:
        return None
    value = beat.get(col)
    if value is None:
        return ""
    return value


def _style_header(ws, ncols: int) -> None:
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(name="Arial", bold=True, color="F7F4EE", size=11)
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}1"
    ws.sheet_properties.tabColor = NAVY


def _write_readme(wb: Workbook, lesson_id: str) -> None:
    ws = wb.active
    ws.title = "Read me"
    lines = [
        "Faculty edit sheet",
        f"Lesson {lesson_id}. The YAML file remains the source. This workbook is a copy for editing.",
        "Edit yellow cells. Do not rename column headers or the beat_id. Do not add a sheet.",
        "Speech sits in en / zh / zh_hk. If a card reveals in steps, use s1, s2, s3 instead of the single en column.",
        "words_en and chars_zh count as you type. They are not imported.",
        "Empty glossary or quiz sheets are ignored. A filled sheet replaces that block.",
        "Beats missing from the Beats sheet are kept. A new beat_id is appended.",
        "Import writes a .bak once, then refuses the save if the lesson no longer validates.",
    ]
    ws["A1"] = lines[0]
    ws["A1"].font = Font(name="Arial", bold=True, size=16, color=NAVY)
    for i, line in enumerate(lines[1:], start=3):
        ws.cell(i, 1, line).font = Font(name="Arial", size=12, color=INK)
        ws.cell(i, 1).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 110
    ws.sheet_properties.tabColor = GOLD


def _write_lesson(wb: Workbook, lesson: dict) -> None:
    ws = wb.create_sheet("Lesson")
    headers = ["field", "en", "zh", "zh_hk"]
    for col, name in enumerate(headers, start=1):
        ws.cell(1, col, name)
    rows = [
        ("title", lesson.get("title_en"), lesson.get("title_zh"), lesson.get("title_zh_hk")),
        ("objective", lesson.get("objective_en"), lesson.get("objective_zh"), lesson.get("objective_zh_hk")),
        ("course", lesson.get("course_en") or lesson.get("course"), lesson.get("course_zh"), lesson.get("course_zh_hk")),
        ("audience", lesson.get("audience_en") or lesson.get("audience"), lesson.get("audience_zh"), lesson.get("audience_zh_hk")),
        ("duration_target_s", lesson.get("duration_target_s"), "", ""),
    ]
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c, "" if value is None else value)
            cell.font = Font(name="Arial", size=11, color=INK)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.number_format = "@"
            if c > 1:
                cell.fill = PatternFill("solid", fgColor=YELLOW)
        ws.row_dimensions[r].height = 36
    _style_header(ws, 4)
    ws.column_dimensions["A"].width = 22
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 56


def _write_beats(wb: Workbook, lesson: dict) -> None:
    ws = wb.create_sheet("Beats")
    beats = lesson.get("beats") or []
    cols = _beat_columns(beats)
    for c, name in enumerate(cols, start=1):
        ws.cell(1, c, name)
    speech_at = {name: cols.index(name) + 1 for name in ("en", "s1_en", "zh", "s1_zh") if name in cols}
    for r, beat in enumerate(beats, start=2):
        for c, name in enumerate(cols, start=1):
            if name == "words_en":
                bits = []
                for src_name in ("s1_en", "en"):
                    src = speech_at.get(src_name)
                    if src:
                        ref = f"{get_column_letter(src)}{r}"
                        bits.append(f'IF({ref}="","",LEN(TRIM({ref}))-LEN(SUBSTITUTE(TRIM({ref})," ",""))+1)')
                cell = ws.cell(r, c, "=" + "+".join(bits) if bits else "")
            elif name == "chars_zh":
                bits = []
                for src_name in ("s1_zh", "zh"):
                    src = speech_at.get(src_name)
                    if src:
                        ref = f"{get_column_letter(src)}{r}"
                        bits.append(f'IF({ref}="","",LEN({ref}))')
                cell = ws.cell(r, c, "=" + "+".join(bits) if bits else "")
            else:
                cell = ws.cell(r, c, _cell_value(beat, name))
                cell.number_format = "@"
            cell.font = Font(name="Arial", size=11, color=INK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if name == "beat_id":
                cell.fill = PatternFill("solid", fgColor=GREY)
            elif name in SPEECH or name.endswith(("_en", "_zh", "_zh_hk")) or name in {"title", "formula"}:
                cell.fill = PatternFill("solid", fgColor=YELLOW)
        ws.row_dimensions[r].height = 48
    _style_header(ws, len(cols))
    widths = {"beat_id": 16, "visual": 18, "seconds": 12, "words_en": 12, "chars_zh": 12}
    for c, name in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 28 if name in SPEECH or name.endswith("_en") else 18)
    if beats:
        dv = DataValidation(type="list", formula1='"' + ",".join(VISUALS) + '"', allow_blank=False)
        dv.error = "Pick a card type the factory knows."
        dv.errorTitle = "Visual"
        ws.add_data_validation(dv)
        dv.add(f"B2:B{len(beats)+1}")


def _write_glossary(wb: Workbook, lesson: dict) -> None:
    ws = wb.create_sheet("Glossary")
    headers = ["en", "zh", "zh_hk", "lock"]
    for c, name in enumerate(headers, start=1):
        ws.cell(1, c, name)
    for r, row in enumerate(lesson.get("glossary") or [], start=2):
        values = [row.get("en"), row.get("zh"), row.get("zh_hk"), "yes" if row.get("lock") else ""]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, "" if value is None else value)
            cell.font = Font(name="Arial", size=11)
            cell.fill = PatternFill("solid", fgColor=YELLOW)
            cell.number_format = "@"
    _style_header(ws, 4)
    for col, width in zip("ABCD", (28, 28, 28, 12)):
        ws.column_dimensions[col].width = width


def _write_quiz(wb: Workbook, lesson: dict) -> None:
    ws = wb.create_sheet("Quiz")
    headers = ["id", "type", "en", "zh", "zh_hk", "answer_en", "answer_zh", "answer_zh_hk"]
    for c, name in enumerate(headers, start=1):
        ws.cell(1, c, name)
    for r, row in enumerate(lesson.get("quiz") or [], start=2):
        values = [row.get("id"), row.get("type"), row.get("en"), row.get("zh"), row.get("zh_hk"), row.get("answer_en"), row.get("answer_zh"), row.get("answer_zh_hk")]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, "" if value is None else value)
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = PatternFill("solid", fgColor=YELLOW)
            cell.number_format = "@"
        ws.row_dimensions[r].height = 48
    _style_header(ws, len(headers))
    for c, width in enumerate([12, 16, 42, 42, 42, 42, 42, 42], start=1):
        ws.column_dimensions[get_column_letter(c)].width = width


def export_lesson(lesson: dict, dest: Path) -> Path:
    wb = Workbook()
    _write_readme(wb, str(lesson.get("id") or "lesson"))
    _write_lesson(wb, lesson)
    _write_beats(wb, lesson)
    _write_glossary(wb, lesson)
    _write_quiz(wb, lesson)
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def _rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [_text(c) for c in rows[0]]
    out = []
    for raw in rows[1:]:
        if not any(_text(c) for c in raw):
            continue
        out.append({header[i]: raw[i] for i in range(min(len(header), len(raw))) if header[i]})
    return out


def _put(beat: dict, key: str, value) -> None:
    text = _text(value)
    if text == "":
        beat.pop(key, None)
        return
    if key == "seconds":
        beat[key] = float(text)
        return
    if key in {"rows", "cols"}:
        beat[key] = int(float(text))
        return
    beat[key] = text


def merge_lesson(lesson: dict, book: Path) -> tuple[dict, list[str]]:
    wb = load_workbook(book, data_only=False)
    notes: list[str] = []
    if "Lesson" in wb.sheetnames:
        for row in _rows(wb["Lesson"]):
            field = _text(row.get("field"))
            if field == "duration_target_s" and _text(row.get("en")):
                lesson["duration_target_s"] = int(float(row["en"]))
            elif field in {"title", "objective", "course", "audience"}:
                for lang, suffix in (("en", "_en"), ("zh", "_zh"), ("zh_hk", "_zh_hk")):
                    if lang in row:
                        _put(lesson, f"{field}{suffix}", row.get(lang))
    if "Beats" in wb.sheetnames:
        by_id = {str(b.get("id")): b for b in lesson.get("beats") or []}
        order = [str(b.get("id")) for b in lesson.get("beats") or []]
        seen = []
        for row in _rows(wb["Beats"]):
            bid = _text(row.get("beat_id"))
            if not bid:
                notes.append("Skipped a Beats row with no beat_id.")
                continue
            beat = by_id.get(bid)
            if beat is None:
                beat = {"id": bid, "visual": _text(row.get("visual")) or "title_card"}
                by_id[bid] = beat
                order.append(bid)
                notes.append(f"Appended beat {bid}.")
            seen.append(bid)
            for key, value in row.items():
                if key in SKIP_IMPORT or key == "beat_id" or not key:
                    continue
                _put(beat, key, value)
            beat["id"] = bid
        lesson["beats"] = [by_id[i] for i in order if i in by_id]
        missing = [i for i in order if i not in seen and i in {str(b.get("id")) for b in lesson["beats"]}]
        if missing:
            notes.append("Kept beats not on the sheet: " + ", ".join(missing))
    if "Glossary" in wb.sheetnames:
        rows = _rows(wb["Glossary"])
        if rows:
            lesson["glossary"] = [
                {
                    "en": _text(r.get("en")),
                    "zh": _text(r.get("zh")),
                    "zh_hk": _text(r.get("zh_hk")),
                    "lock": _text(r.get("lock")).lower() in {"1", "yes", "true", "y", "lock"},
                }
                for r in rows if _text(r.get("en"))
            ]
    if "Quiz" in wb.sheetnames:
        rows = _rows(wb["Quiz"])
        if rows:
            lesson["quiz"] = [
                {
                    "id": _text(r.get("id")) or f"Q{i}",
                    "type": _text(r.get("type")) or "comprehension",
                    "en": _text(r.get("en")),
                    "zh": _text(r.get("zh")),
                    "zh_hk": _text(r.get("zh_hk")),
                    "answer_en": _text(r.get("answer_en")),
                    "answer_zh": _text(r.get("answer_zh")),
                    "answer_zh_hk": _text(r.get("answer_zh_hk")),
                }
                for i, r in enumerate(rows, start=1)
                if _text(r.get("en"))
            ]
    return lesson, notes


def import_workbook(yaml_path: Path, book: Path, dump) -> dict:
    lesson = load_lesson(yaml_path)
    merged, notes = merge_lesson(lesson, book)
    errors = validate(merged, glyphs=False)
    if errors:
        return {"ok": False, "error": "Import refused. The sheet no longer validates.", "errors": errors, "notes": notes}
    bak = yaml_path.with_suffix(yaml_path.suffix + ".bak")
    if not bak.exists():
        bak.write_text(yaml_path.read_text(encoding="utf-8"), encoding="utf-8")
    yaml_path.write_text(dump(merged), encoding="utf-8")
    return {"ok": True, "file": yaml_path.name, "backup": bak.name, "notes": notes}


def main() -> int:
    if len(sys.argv) < 3 or sys.argv[1] not in {"export", "import"}:
        print("Usage: xlsx_lesson.py export LESSON.yaml [out.xlsx]", file=sys.stderr)
        print("       xlsx_lesson.py import LESSON.yaml sheet.xlsx", file=sys.stderr)
        return 2
    src = Path(sys.argv[2])
    if sys.argv[1] == "export":
        lesson = load_lesson(src)
        dest = Path(sys.argv[3]) if len(sys.argv) > 3 else src.with_suffix(".xlsx")
        export_lesson(lesson, dest)
        print(dest)
        return 0
    if len(sys.argv) < 4:
        print("import needs the xlsx path", file=sys.stderr)
        return 2
    from app import _dump_yaml
    result = import_workbook(src, Path(sys.argv[3]), _dump_yaml)
    print(result)
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
